"""Real offline-conversion and hostile-input regressions; no public service used."""
import io
import json
from pathlib import Path
import socket
import subprocess
import sys
import tempfile
import unittest
import zipfile

WORKER = Path(__file__).with_name("convert.py")


def docx(document_xml=None, relationship=None):
    document_xml = document_xml or '''<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body><w:p><w:r><w:t>中文 Hello document</w:t></w:r></w:p></w:body></w:document>'''
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", '''<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"><Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/><Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/></Types>''')
        archive.writestr("_rels/.rels", '''<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/></Relationships>''')
        archive.writestr("word/document.xml", document_xml)
        if relationship:
            archive.writestr("word/_rels/document.xml.rels", relationship)
    return buffer.getvalue()


def pdf():
    text = b"BT /F1 12 Tf 50 750 Td (MarkItDown PDF receipt) Tj ET"
    objects = [b"<< /Type /Catalog /Pages 2 0 R >>",
               b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
               b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
               b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
               b"<< /Length " + str(len(text)).encode() + b" >>\nstream\n" + text + b"\nendstream"]
    data = b"%PDF-1.4\n"
    offsets = [0]
    for index, obj in enumerate(objects, 1):
        offsets.append(len(data))
        data += str(index).encode() + b" 0 obj\n" + obj + b"\nendobj\n"
    start = len(data)
    data += b"xref\n0 6\n0000000000 65535 f \n"
    for offset in offsets[1:]:
        data += f"{offset:010d} 00000 n \n".encode()
    return data + b"trailer\n<< /Size 6 /Root 1 0 R >>\nstartxref\n" + str(start).encode() + b"\n%%EOF"


class ConversionTests(unittest.TestCase):
    def run_worker(self, extension, contents):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / ("private filename" + extension)
            path.write_bytes(contents.encode("utf-8") if isinstance(contents, str) else contents)
            result = subprocess.run([sys.executable, "-I", str(WORKER), str(path)],
                                    capture_output=True, timeout=45)
            self.assertEqual(result.stderr, b"")
            response = json.loads(result.stdout.decode("utf-8"))
            self.assertNotIn(str(path), result.stdout.decode("utf-8"))
            return result.returncode, response

    def assert_markdown(self, extension, contents, expected):
        code, response = self.run_worker(extension, contents)
        self.assertEqual(code, 0, response)
        self.assertEqual(response["version"], "0.1.7")
        self.assertIn(expected, response["markdown"])

    def test_docx_real_converter(self):
        self.assert_markdown(".docx", docx(), "中文 Hello document")

    def test_pdf_real_converter(self):
        self.assert_markdown(".pdf", pdf(), "MarkItDown PDF receipt")

    def test_plain_text_and_markdown_unicode(self):
        for extension in (".txt", ".md"):
            self.assert_markdown(extension, "# Hello\n中文 日本語 café", "中文 日本語 café")

    def test_html_formatting_and_no_script(self):
        code, response = self.run_worker(".html", "<h1>Title</h1><p><strong>Bold</strong></p><script>PRIVATE_SCRIPT</script>")
        self.assertEqual(code, 0, response)
        self.assertIn("**Bold**", response["markdown"])
        self.assertNotIn("PRIVATE_SCRIPT", response["markdown"])

    def test_html_remote_resources_never_requested(self):
        with socket.socket() as listener:
            listener.bind(("127.0.0.1", 0))
            listener.listen()
            listener.settimeout(0.1)
            port = listener.getsockname()[1]
            html = f'<html><head><link href="http://127.0.0.1:{port}/style.css" rel="stylesheet"></head><body><h1>Offline</h1><img src="http://127.0.0.1:{port}/secret"><iframe src="https://www.youtube.com/watch?v=example"></iframe><a href="https://www.youtube.com/watch?v=example">Video link</a></body></html>'
            self.assert_markdown(".html", html, "Offline")
            with self.assertRaises(TimeoutError):
                listener.accept()

    def test_csv_table(self):
        self.assert_markdown(".csv", "Name,Value\nAlice,12\n", "| Alice | 12 |")

    def test_xlsx_real_converter(self):
        import openpyxl
        buffer = io.BytesIO()
        workbook = openpyxl.Workbook()
        workbook.active.append(["Name", "Value"])
        workbook.active.append(["Alice", 42])
        workbook.save(buffer)
        self.assert_markdown(".xlsx", buffer.getvalue(), "Alice")

    def test_pptx_real_converter(self):
        from pptx import Presentation
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[0])
        slide.shapes.title.text = "Presentation import"
        buffer = io.BytesIO()
        presentation.save(buffer)
        self.assert_markdown(".pptx", buffer.getvalue(), "Presentation import")

    def test_entities_refused_before_conversion(self):
        document = '<!DOCTYPE x [<!ENTITY private SYSTEM "file:///private/etc/passwd">]><x>&private;</x>'
        code, response = self.run_worker(".docx", docx(document))
        self.assertNotEqual(code, 0)
        self.assertEqual(response, {"error": "unsafe_document"})

    def test_utf16_entities_refused(self):
        document = '<?xml version="1.0" encoding="UTF-16"?><!DOCTYPE x [<!ENTITY private SYSTEM "file:///private/etc/passwd">]><x>&private;</x>'
        code, response = self.run_worker(".docx", docx(document.encode("utf-16")))
        self.assertNotEqual(code, 0)
        self.assertEqual(response, {"error": "unsafe_document"})

    def test_office_external_image_refused(self):
        relationship = '<Relationships><Relationship Type="example/image" TargetMode="External" Target="file:///private/secret"/></Relationships>'
        code, response = self.run_worker(".docx", docx(relationship=relationship))
        self.assertNotEqual(code, 0)
        self.assertEqual(response, {"error": "external_resource"})

    def test_hyperlinks_remain_supported(self):
        relationship = '<Relationships><Relationship Type="example/hyperlink" TargetMode="External" Target="https://example.com"/></Relationships>'
        self.assert_markdown(".docx", docx(relationship=relationship), "Hello document")

    def test_zip_bomb_refused(self):
        buffer = io.BytesIO()
        with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
            archive.writestr("word/document.xml", b"x" * 1_000_000)
        code, response = self.run_worker(".docx", buffer.getvalue())
        self.assertNotEqual(code, 0)
        self.assertEqual(response, {"error": "archive_limits"})

    def test_entry_limit_refused(self):
        buffer = io.BytesIO()
        with zipfile.ZipFile(buffer, "w") as archive:
            for index in range(2049):
                archive.writestr(str(index), b"")
        code, response = self.run_worker(".docx", buffer.getvalue())
        self.assertNotEqual(code, 0)
        self.assertEqual(response, {"error": "archive_limits"})

    def test_input_limit(self):
        code, response = self.run_worker(".txt", b"x" * (25 * 1024 * 1024 + 1))
        self.assertNotEqual(code, 0)
        self.assertEqual(response, {"error": "input_too_large"})

    def test_output_limit(self):
        code, response = self.run_worker(".txt", b"x" * 2_000_001)
        self.assertNotEqual(code, 0)
        self.assertEqual(response, {"error": "output_too_large"})

    def test_unsupported_and_empty(self):
        self.assertEqual(self.run_worker(".exe", b"example")[1], {"error": "unsupported_format"})
        self.assertEqual(self.run_worker(".txt", b"")[1], {"error": "empty_document"})

    def test_damaged_document_has_no_private_diagnostics(self):
        code, response = self.run_worker(".pdf", b"PRIVATE document malformed PDF")
        self.assertNotEqual(code, 0)
        self.assertEqual(response, {"error": "conversion_failed"})

    def test_runtime_guard_denies_network_subprocess_and_external_reads(self):
        with tempfile.TemporaryDirectory() as directory:
            secret = Path(directory) / "private.txt"
            secret.write_text("PRIVATE", encoding="utf-8")
            script = f'''
import runpy, socket, subprocess, sys
worker = runpy.run_path({str(WORKER)!r})
worker["restrict_io"]()
assert isinstance(socket.gethostname(), str)
actions = [lambda: socket.socket(), lambda: subprocess.run([sys.executable, "-V"]),
           lambda: open({str(secret)!r}).read(), lambda: open(sys.prefix + "/forbidden.txt", "w")]
for action in actions:
    try:
        action()
        raise AssertionError("action escaped guard")
    except worker["ImportFailure"] as error:
        assert error.code == "external_resource"
print("guard-denied-4")
'''
            result = subprocess.run([sys.executable, "-I", "-c", script], capture_output=True, timeout=20)
            self.assertEqual(result.returncode, 0, result.stderr.decode())
            self.assertEqual(result.stdout.strip(), b"guard-denied-4")


if __name__ == "__main__":
    unittest.main()
